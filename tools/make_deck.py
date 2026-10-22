#!/usr/bin/env python3
"""生成《潜影》设计展示 PPT 初版（16 页，讲稿在备注栏）。
用法：python3 tools/make_deck.py   →  docs/展示/潜影-设计展示-初版.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "docs/图"
SHOT = IMG / "实机截图-10-14"
OUT = ROOT / "docs/展示/潜影-设计展示-初版.pptx"

BG, FG, DIM, ACC = RGBColor(0x14, 0x12, 0x10), RGBColor(0xEE, 0xE8, 0xDC), RGBColor(0x9A, 0x92, 0x85), RGBColor(0xC8, 0x4B, 0x3C)
FONT = "PingFang SC"

# 每页：(类型, 标题, 要点[(文字, 缩进)], 图[文件或"占位:说明"], 讲稿)
S = []
def add(kind, title, pts=(), imgs=(), notes=""):
    S.append((kind, title, list(pts), list(imgs), notes))

add("cover", "潜影", [("Latent Image　·　奥瑞尔庄园 · 一九二六", 0), ("第一人称推理叙事游戏", 0),
    ("第十四届全国大学生数字媒体科技作品及创意竞赛 · 虚拟现实与游戏类", 0), ("队名 / 成员：【待填】", 0)],
    [IMG / "滤镜-开场.png"], "封面。建议换成红灯下显影盘里浮出一张脸的画面（脸看不清）。")
add("body", "目录", [("01　作品是什么　故事与世界观", 0), ("02　怎么玩　一天的循环与高潮", 0),
    ("03　创新在哪　四个创新点", 0), ("04　做到了什么　工程与成果", 0)], [], "")
add("section", "01　作品是什么", [("一张已经曝光、还没冲洗的底片", 0)], [], "这一句就是“潜影”的意思，也是接下来三页要讲清楚的东西。")
add("body", "一句话定位与命名", [
    ("1926 年秋，侦探应老友之邀来到庄园，当晚主人死在自己反锁的暗房里。最后，侦探发现凶手是他自己。", 0),
    ("第一人称 3D 推理叙事游戏｜桌面（macOS / Windows）｜Godot 4｜一周目 20—30 分钟", 0),
    ("为什么叫“潜影”", 0), ("案件：被忘掉的星期四夜里", 1), ("怪谈：合影里还没出现的脸", 1), ("玩法：玩家亲手把它冲洗出来", 1)],
    ["占位:庄园外景，叠标题字"], "作品名一语多关，好记、好搜，也不剧透。副标题交代年代地点。灵媒“拉坦特”（Latente）的名字也和它呼应。")
add("body", "故事与怪谈", [
    ("案件：主人死在两道门都从里面闩上的暗房；验尸官一周后开庭", 0),
    ("怪谈：每年全家合影里多出一张谁也不认识的脸；第二年，那张脸的主人被所有人忘记", 0),
    ("六个人物：侦探、主人、女主人、孩子、园丁、灵媒", 0),
    ("八章鱼骨结构：序章 → 五章各围绕一个人 → 破案 → 镜子与尾声", 0),
    ("每章结束嫌疑落在一个人身上，下一章推翻", 0)],
    ["占位:六个人物模型并排（docs/人物/）；八章标题连成的诗"], "这一页只讲前提，不剧透到底。强调每章一个嫌疑人、一个推论，玩家总有明确目标。")
add("body", "三层叙事与想表达的东西", [
    ("L3 故事：1926 年的庄园（从头就在看）", 0), ("L1 容器：侦探是数字容器里的一个人格，还有别的“我”", 0),
    ("L0 世界：一位病人的不同人格被放进容器（尾声揭示，不解释机构）", 0),
    ("追问：人如何对待由自己分裂出来的那些“自己”", 0),
    ("落点是正向的：不把精神疾病与暴力挂钩，“杀人”在上层的含义是遗忘", 0)],
    [IMG / "滤镜-L0胶片.png", IMG / "滤镜-L3录像带.png", IMG / "滤镜-L4数字故障.png"],
    "世界是黑的，落点是正的。玩家一层层往外走，不是被告知。这一页要把“不做什么”的底线讲一句。")
add("section", "02　怎么玩", [("走进现场，把线索连成一个结论", 0)], [], "")
add("body", "一天的循环", [
    ("客房是玩家的桌面：床、办公桌、打字机、记事本", 0),
    ("白天：自由走动 → 一对一问话 → E 键检视物件 → 线索进记事本", 0),
    ("傍晚：证据板连线得出本章推论；连错时侦探自言自语，提示该去找哪些线索", 0),
    ("夜里：打字机写当天记录 → 睡下 → 梦里再过一遍", 0),
    ("部分线索可深入探查，进入独立的演绎关卡（密室、合影、回廊、星期四夜里）", 0)],
    [SHOT / "2-第一二章.jpg"], "五章五个推论，节奏一致，新手不会迷路。强调“线索藏在空间和物件里，不藏在对话框里”。截图是 10-14 的，建议换最新。")
add("body", "空间就是证据", [
    ("完整搭建的庄园：门厅、餐厅、客厅、温室、档案室、暗房、白色走廊……", 0),
    ("体验准则：P.T. 式的空间感，重访即变化，细思极恐而不突脸", 0),
    ("世界观靠“接缝”渗漏：L0 胶片 → L3 录像带 → L4 数字故障，没有人解释", 0),
    ("画面做法：先调灯光与天空，再调材质，最后才调滤镜（基准：《杀手 3》达特穆尔）", 0)],
    ["占位:门厅同一机位 L0 / L3 / L4 三联图（视频镜头 S20a/b/c）"], "同一个房间在不同渗漏等级下的三种样子是这一页的重点。")
add("body", "高潮：亲手冲洗底片，与两种讲法", [
    ("第 6 章：玩家在红灯下亲手冲洗叠印底片，多出来的脸慢慢浮上来，是侦探自己", 0),
    ("第 7 章：镜子里一直照不出他。玩家选择相信哪一种讲法：", 0),
    ("为了保护所有人", 1), ("为了不被取代", 1), ("开发者不裁决", 0)],
    [SHOT / "4-终局与尾声.jpg"], "答案不是看过场，而是在玩家手里浮现。两种讲法借鉴《少年派》。建议换成显影盘画面（S04）。")
add("section", "03　创新在哪", [("四个创新点", 0)], [], "")
add("body", "创新点 1 · 2：叙事与意象", [
    ("① 第一人称的叙述诡计", 0), ("参照《罗杰疑案》：叙述者就是凶手，只省略、不说谎", 1),
    ("省略落在镜头和空间：镜头避开倒影、对话从不同时发生、旁白句句真话但永远不说“第一次来”", 1),
    ("伏笔埋在共有的旧疤、敲桌习惯和笔迹里", 1),
    ("② 玩法与摄影意象同构", 0), ("潜影（发生了却看不见）→ 显影（想起来）→ 定影（只留下一个）", 1),
    ("同时是案件结构、怪谈规律、人格隐喻，也是玩家的关键动作", 1)],
    ["占位:倒影缺席的镜头截图；三步工艺示意图"], "先把“让读者亲手揭穿自己”讲出来，再讲玩法和意象为什么是同一件事。这两条是最扎实的创新，因为游戏里现在就有。")
add("body", "创新点 3 · 4：共享记录与工程管线", [
    ("③ 异步共享的记录（设计，待实现）", 0),
    ("世界观里本来就有十二个容器、十二个“格雷”", 1),
    ("玩家写下的记录留存，后来的玩家在终章读到前人那一页，再为下一个人留一句话", 1),
    ("只用玩家自己写过的预设句子，不开放自由输入；借鉴《尼尔：机械纪元》", 1),
    ("④ 分镜即数据", 0), ("剧情、镜头、台词、证据、音效全在表格里，代码不写死内容", 1),
    ("内容校验、两种讲法自动通关、截图与帧率检测", 1),
    ("一人加 coding agent 撑起完整生产线，队友只交表格和文件", 1)],
    ["占位:左 十二块屏幕的尾声；右 CSV 表格与自动通关终端截图"],
    "对应评审的“多人互动体验”和“技术实现、运行稳定性”。【核实】异步共享记录在 docs 和代码里没有找到落地记录，若未完成，请保持“设计”的标注，不要写成已实现。")
add("body", "工程与真实成果", [
    ("管线：Blender 建模烘焙 → Godot 关卡；玩家本体八方向实测碰撞，防穿模", 0),
    ("镜头 85 / 台词 94 / 证据 55；6 个现行人物；八章六组推论；两种终局（以《检查报告.md》为准）", 0),
    ("验证：内容校验 + 两种讲法自动通关都通过；重点场景不低于 45 帧", 0),
    ("AI：AI 原生开发（代码由 AI 编程助手完成，设计与审美判断由我们负责）；“千人千面”目前是设计加人写的兜底，游戏里没有运行时调用", 0)],
    ["占位:Blender 线框 → Godot 实时对比，或数据卡"], "这一页讲真实的东西。AI 那一条照实讲，不演示实时生成。数字来自 10-22 的视频脚本，上屏前再核一遍检查报告。")
add("body", "创意来源与应用价值", [
    ("创意来源", 0), ("《罗杰疑案》：叙述者即凶手", 1), ("博尔赫斯：《死亡与罗盘》《环形废墟》《特隆、乌克巴尔、奥比斯·特蒂乌斯》《巴别图书馆》", 1),
    ("《少年派》：两种讲法　《P.T.》：回廊与重访　《明日方舟》“黑流树海”PV：做旧录像带质感", 1),
    ("摄影史：塔尔博特的第一张负片（庄园名“奥瑞尔”、死者姓“塔尔博特”均出于此）", 1),
    ("应用价值", 0), ("文化与历史：灵魂照片、唯灵论、早期摄影工艺", 1),
    ("互动叙事方法：把叙述诡计和分层叙事搬进 3D 空间", 1), ("小团队开发：表格驱动加 AI 编程助手，不写代码的队友也能做内容", 1)],
    [], "来源别念全表，挑能解释设计的几条。“奥瑞尔”与“塔尔博特”出自第一张负片那扇凸窗，是个好彩蛋。")
add("end", "杀人是遗忘，调查是记起。", [("《潜影》 Latent Image", 0), ("团队 · 联系方式 · 致谢（队友、第三方素材来源）：【待填】", 0)],
    [SHOT / "4-终局与尾声.jpg"], "收一句“两位玩家，两次游玩，过程不同，真相相同”，然后停住，留 Q&A。建议换尾声十二块屏幕（S25）。")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def text(slide, x, y, w, h, pts, size, color=FG, bold_first=False, align=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, (t, lv) in enumerate(pts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("　　" * lv) + ("– " if lv else "") + t
        p.space_after = Pt(8)
        if align: p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size - 3 * lv); r.font.name = FONT
            r.font.color.rgb = color if lv == 0 else DIM
    return tb

def pic_or_box(slide, spec, x, y, w, h):
    if isinstance(spec, Path) and spec.exists():
        pic = slide.shapes.add_picture(str(spec), Inches(x), Inches(y))
        r = min(w / (pic.width / 914400), h / (pic.height / 914400))
        pic.width, pic.height = int(pic.width * r), int(pic.height * r)
        pic.left, pic.top = Inches(x) + (Inches(w) - pic.width) // 2, Inches(y)
    else:
        label = spec if isinstance(spec, str) else f"缺图：{spec.name}"
        box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x2A, 0x26, 0x22); box.line.color.rgb = DIM
        tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.text = "配图：" + label.replace("占位:", "")
        for r in tf.paragraphs[0].runs: r.font.size = Pt(14); r.font.name = FONT; r.font.color.rgb = DIM
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

for n, (kind, title, pts, imgs, notes) in enumerate(S, 1):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    if kind in ("section", "cover", "end"):
        big = 66 if kind != "end" else 44
        text(s, 0.9, 2.2 if kind != "cover" else 1.6, 11.5, 1.6, [(title, 0)], big, ACC if kind == "section" else FG)
        text(s, 0.9, 3.9 if kind != "cover" else 3.3, 7.5 if imgs else 11.5, 3, pts, 22, DIM if kind == "section" else FG)
        if imgs and kind != "section": pic_or_box(s, imgs[0], 8.6, 1.4, 4.2, 4.6)
    else:
        text(s, 0.7, 0.45, 12, 1, [(title, 0)], 34, ACC)
        three = len(imgs) == 3
        wide = bool(imgs) and not three
        text(s, 0.7, 1.5, 7.3 if wide else (12 if not imgs else 12), 4.2 if three else 5.6, pts, 20)
        if wide: pic_or_box(s, imgs[0], 8.3, 1.6, 4.6, 4.8)
        if three:
            for i, im in enumerate(imgs): pic_or_box(s, im, 0.7 + i * 4.1, 5.0, 3.9, 2.1)
    text(s, 12.2, 7.0, 1, 0.4, [(str(n), 0)], 12, DIM, align=PP_ALIGN.RIGHT)
    if notes: s.notes_slide.notes_text_frame.text = notes

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print("saved", OUT, len(S), "slides")
